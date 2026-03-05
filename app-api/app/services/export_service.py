from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from app.core.settings import ROOT_DIR, load_app_config, load_export_columns
from app.models import ExamSummaryRecord


RED_FONT = Font(color="FF0000", bold=True)


DEFAULT_COLUMNS = [
    {"field": "org_name", "title": "单位"},
    {"field": "person_name", "title": "姓名"},
    {"field": "gender", "title": "性别"},
    {"field": "id_no", "title": "证件号"},
    {"field": "phone", "title": "电话"},
    {"field": "exam_no", "title": "体检编号"},
    {"field": "summary_date", "title": "汇总日期"},
    {"field": "final_date", "title": "终检日期"},
    {"field": "exam_status", "title": "体检状态"},
    {"field": "group_name", "title": "项目组"},
    {"field": "item_name", "title": "参数名称"},
    {"field": "result_value", "title": "结果值"},
    {"field": "unit", "title": "单位"},
    {"field": "ref_range", "title": "参考范围"},
    {"field": "abnormal_flag", "title": "异常标记"},
]


def _get_columns() -> list[dict[str, str]]:
    configured = load_export_columns()
    return configured if configured else DEFAULT_COLUMNS


def _resolve_export_dir(requested_export_dir: str | None) -> Path:
    app_cfg = load_app_config()
    if requested_export_dir and requested_export_dir.strip():
        target = Path(requested_export_dir.strip())
    else:
        target = Path(app_cfg.export_dir)

    if not target.is_absolute():
        target = ROOT_DIR / target
    target.mkdir(parents=True, exist_ok=True)
    return target


def _allow_group(group_name: str, selected_groups: set[str] | None) -> bool:
    if not selected_groups:
        return True
    return group_name in selected_groups


def export_records(
    records: list[ExamSummaryRecord],
    export_dir: str | None = None,
    selected_groups: list[str] | None = None,
) -> tuple[str, str]:
    target_dir = _resolve_export_dir(export_dir)
    selected_group_set = set(selected_groups) if selected_groups else None

    file_name = f"体检导出_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    file_path = target_dir / file_name

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "体检数据"

    columns = _get_columns()
    headers = [column["title"] for column in columns]
    sheet.append(headers)

    for index, title in enumerate(headers, start=1):
        sheet.cell(row=1, column=index).font = Font(bold=True)
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(len(title) * 2, 12)

    for record in records:
        for group in record.project_groups:
            if not _allow_group(group.group_name, selected_group_set):
                continue

            if not group.items:
                row_data = {
                    "org_name": record.org_name,
                    "person_name": record.person_name,
                    "gender": record.gender,
                    "id_no": record.id_no,
                    "phone": record.phone,
                    "exam_no": record.exam_no,
                    "summary_date": record.summary_date,
                    "final_date": record.final_date,
                    "exam_status": record.exam_status,
                    "group_name": group.group_name,
                    "item_name": "",
                    "result_value": "",
                    "unit": "",
                    "ref_range": "",
                    "abnormal_flag": "",
                }
                row = [str(row_data.get(column["field"], "") or "") for column in columns]
                sheet.append(row)
                continue

            for item in group.items:
                row_data = {
                    "org_name": record.org_name,
                    "person_name": record.person_name,
                    "gender": record.gender,
                    "id_no": record.id_no,
                    "phone": record.phone,
                    "exam_no": record.exam_no,
                    "summary_date": record.summary_date,
                    "final_date": record.final_date,
                    "exam_status": record.exam_status,
                    "group_name": group.group_name,
                    "item_name": item.item_name,
                    "result_value": item.result_value,
                    "unit": item.unit,
                    "ref_range": item.ref_range,
                    "abnormal_flag": item.abnormal_flag,
                }
                row = [str(row_data.get(column["field"], "") or "") for column in columns]
                sheet.append(row)

                if item.is_abnormal:
                    row_index = sheet.max_row
                    for column_index, column in enumerate(columns, start=1):
                        if column["field"] in {"result_value", "abnormal_flag", "item_name"}:
                            sheet.cell(row=row_index, column=column_index).font = RED_FONT

    workbook.save(file_path)
    return file_name, str(file_path)
